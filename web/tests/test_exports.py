"""Export API tests read actual workbooks, with no SSH or target-site requests."""

import copy
import importlib
import os
import unittest
from io import BytesIO
from unittest.mock import patch

import exports
from openpyxl import load_workbook

RECORD = {
    "query": "실리콘 주걱",
    "items": [{"name": "주걱", "price": 0, "reviews": 12, "fee": 0, "ad": False}],
    "related": ["요리 주걱"],
    "autocomplete": ["실리콘 주걱 세트"],
    "sections": {"search": "ok", "autocomplete": "ok", "sourcing": "ok"},
    "dome_exists": False,
    "dome_count": 0,
    "opportunity": 0.75,
}


class ExportAPI(unittest.TestCase):
    def setUp(self):
        import app

        self.env = patch.dict(os.environ, {"APP_PASSWORD": "export-test"})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.module = importlib.reload(app)
        self.client = self.module.app.test_client()
        self.headers = {"X-App-Password": "export-test"}

    def post(self, records):
        return self.client.post("/api/export", json={"records": records}, headers=self.headers)

    def test_xlsx_contains_all_sections_and_numeric_zero_without_worker_calls(self):
        with patch.object(self.module.miniclient, "worker_request") as worker:
            response = self.post([RECORD])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertIn(".xlsx", response.headers["Content-Disposition"])
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(book.sheetnames, ["키워드 요약", "상품", "연관검색어", "자동완성", "도매꾹"])
        self.assertEqual(book["상품"]["A2"].value, "실리콘 주걱")
        self.assertEqual(book["상품"]["D2"].value, 0)
        self.assertEqual(book["상품"]["D2"].data_type, "n")
        self.assertEqual(book["연관검색어"]["C2"].value, "요리 주걱")
        self.assertEqual(book["자동완성"]["C2"].value, "실리콘 주걱 세트")
        self.assertEqual(book["도매꾹"]["C2"].value, "없음")
        self.assertEqual(book["도매꾹"]["D2"].value, 0)
        worker.assert_not_called()
        book.close()

    def test_selected_payload_does_not_include_unselected_keywords(self):
        other = {**RECORD, "query": "다른 키워드"}
        response = self.post([other])
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(book["키워드 요약"].max_row, 2)
        self.assertEqual(book["상품"]["A2"].value, "다른 키워드")
        book.close()

    def test_failures_are_not_exported_as_zero_or_no_results(self):
        data = {**RECORD, "sections": {"search": "pending", "autocomplete": "http_error", "sourcing": "timeout"}}
        book = load_workbook(BytesIO(self.post([data]).data))
        self.assertEqual(
            [book["키워드 요약"].cell(2, i).value for i in (2, 3, 4)], ["조회 중", "조회 실패", "시간 초과"]
        )
        self.assertEqual(book["도매꾹"]["C2"].value, "미확인")
        self.assertIsNone(book["도매꾹"]["D2"].value)
        book.close()

    def test_formula_names_remain_literal_and_xml_controls_are_removed(self):
        data = copy.deepcopy(RECORD)
        data["items"][0]["name"] = '=HYPERLINK("https://example.test")\x00'
        data["related"] = ["=1+1", "+2", "@SUM(A1)"]
        book = load_workbook(BytesIO(self.post([data]).data), data_only=False)
        self.assertEqual(book["상품"]["C2"].value, '=HYPERLINK("https://example.test")')
        self.assertEqual(book["상품"]["C2"].data_type, "s")
        self.assertEqual(book["연관검색어"]["C2"].data_type, "s")
        book.close()

    def test_requires_password_and_accepts_existing_json_password_contract(self):
        self.assertEqual(self.client.post("/api/export", json={"records": [RECORD]}).status_code, 401)
        response = self.client.post("/api/export", json={"records": [RECORD], "password": "export-test"})
        self.assertEqual(response.status_code, 200)

    def test_record_count_boundaries(self):
        for count, expected in ((0, 400), (1, 200), (499, 200), (500, 200), (501, 400)):
            with self.subTest(count=count):
                self.assertEqual(self.post([{"query": str(i)} for i in range(count)]).status_code, expected)

    def test_collection_size_boundaries(self):
        for field in ("items", "related", "autocomplete"):
            for count, expected in ((199, 200), (200, 200), (201, 400)):
                with self.subTest(field=field, count=count):
                    value = {} if field == "items" else "검색어"
                    self.assertEqual(self.post([{**RECORD, field: [value] * count}]).status_code, expected)

    def test_query_length_boundaries(self):
        for count, expected in ((59, 200), (60, 200), (61, 400)):
            with self.subTest(count=count):
                self.assertEqual(self.post([{**RECORD, "query": "가" * count}]).status_code, expected)

    def test_invalid_payloads_and_duplicates(self):
        for records in (
            None,
            [None],
            [RECORD, RECORD],
            [{"query": " "}],
            [{**RECORD, "items": [None]}],
            [{**RECORD, "related": [{}]}],
            [{**RECORD, "sections": []}],
        ):
            with self.subTest(records=records):
                self.assertEqual(self.post(records).status_code, 400)
        self.assertEqual(self.client.post("/api/export", data="not-json", headers=self.headers).status_code, 400)

    def test_request_body_limit(self):
        # Exercise the route's actual configured byte limit without allocating workbooks.
        for size in (exports.MAX_BODY - 1, exports.MAX_BODY, exports.MAX_BODY + 1):
            with self.subTest(size=size):
                response = self.client.post("/api/export", data=b" " * size, headers=self.headers)
                self.assertEqual(response.status_code, 413 if size > exports.MAX_BODY else 400)


if __name__ == "__main__":
    unittest.main()
