"""Build Excel workbooks from the results already displayed in the browser."""

from __future__ import annotations

import math
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

MAX_RECORDS = 500
MAX_ITEMS = 200
MAX_KEYWORDS = 200
MAX_BODY = 8 * 1024 * 1024
_CONTROLS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_LABELS = {
    "ok": "완료",
    "no_results": "결과 없음",
    "pending": "조회 중",
    "timeout": "시간 초과",
    "http_error": "조회 실패",
    "load_error": "조회 실패",
    "challenge": "접근 제한",
    "paused": "대기 중",
    "budget": "오늘 한도 초과",
    "unknown": "미확인",
}


def _text(value: object) -> str:
    return _CONTROLS.sub("", str(value))[:32767] if value is not None else ""


def _number(value: object) -> int | float | None:
    try:
        return value if type(value) in (int, float) and math.isfinite(value) else None
    except OverflowError:
        return None


def _row(sheet, values: list) -> None:
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(cell.value, str):
            # Keep formula-looking product names and keywords as literal text.
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cells.append(cell)
    sheet.append(cells)


def workbook_bytes(payload: object) -> BytesIO:
    """Validate a bounded selection and return a genuine .xlsx file without any SSH calls."""
    if not isinstance(payload, dict) or not isinstance(payload.get("records"), list):
        raise ValueError("저장할 검색 결과가 필요합니다.")
    records = payload["records"]
    if not 1 <= len(records) <= MAX_RECORDS:
        raise ValueError(f"키워드를 1~{MAX_RECORDS}개 선택해 주세요.")
    seen = set()
    for record in records:
        if not isinstance(record, dict) or not isinstance(record.get("query"), str) or not record["query"].strip():
            raise ValueError("키워드 형식이 올바르지 않습니다.")
        if len(record["query"]) > 60 or record["query"] in seen:
            raise ValueError("키워드가 너무 길거나 중복되었습니다.")
        seen.add(record["query"])
        for field, limit in (("items", MAX_ITEMS), ("related", MAX_KEYWORDS), ("autocomplete", MAX_KEYWORDS)):
            entries = record.get(field, [])
            if not isinstance(entries, list) or len(entries) > limit:
                raise ValueError("저장할 결과의 크기나 형식이 올바르지 않습니다.")
            if any(not isinstance(entry, dict if field == "items" else str) for entry in entries):
                raise ValueError("저장할 항목 형식이 올바르지 않습니다.")
        if not isinstance(record.get("sections", {}), dict):
            raise ValueError("조회 상태 형식이 올바르지 않습니다.")

    book = Workbook()
    summary = book.active
    summary.title = "키워드 요약"
    products = book.create_sheet("상품")
    related = book.create_sheet("연관검색어")
    auto = book.create_sheet("자동완성")
    sourcing = book.create_sheet("도매꾹")
    _row(
        summary,
        [
            "키워드",
            "검색 상태",
            "자동완성 상태",
            "도매꾹 상태",
            "저장 상품 수",
            "총 검색 수",
            "기회점수",
            "경쟁",
            "수요",
            "꾸준함",
            "수집 시각",
            "샘플",
        ],
    )
    _row(products, ["키워드", "순서", "상품명", "가격", "리뷰", "배지", "배송비", "광고", "상품 URL"])
    _row(related, ["키워드", "순서", "연관검색어"])
    _row(auto, ["키워드", "순서", "자동완성"])
    _row(sourcing, ["키워드", "조회 상태", "상품 유무", "상품 수"])
    for record in records:
        query = _text(record["query"])
        sections = record.get("sections", {})
        statuses = [_LABELS.get(_text(sections.get(key)), "미확인") for key in ("search", "autocomplete", "sourcing")]
        items = record.get("items", [])
        _row(
            summary,
            [
                query,
                *statuses,
                len(items),
                _number(record.get("total_count")),
                *[_number(record.get(key)) for key in ("opportunity", "rarity", "demand", "steadiness")],
                _text(record.get("collected_at")),
                "예" if record.get("sample") is True else "아니오",
            ],
        )
        for index, item in enumerate(items, 1):
            _row(
                products,
                [
                    query,
                    index,
                    _text(item.get("name")),
                    _number(item.get("price")),
                    _number(item.get("reviews")),
                    _text(item.get("badge")),
                    _number(item.get("fee")),
                    "예" if item.get("ad") is True else "아니오" if item.get("ad") is False else "",
                    _text(item.get("url")),
                ],
            )
        for sheet, field in ((related, "related"), (auto, "autocomplete")):
            for index, word in enumerate(record.get(field, []), 1):
                _row(sheet, [query, index, _text(word)])
        exists = record.get("dome_exists")
        valid_source = sections.get("sourcing") == "ok"
        _row(
            sourcing,
            [
                query,
                statuses[2],
                ("있음" if exists else "없음") if valid_source and exists is not None else "미확인",
                _number(record.get("dome_count")) if valid_source else None,
            ],
        )

    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(14, width * 1.3))
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output
